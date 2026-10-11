import sys
sys.path.insert(0, "/home/claude/adapt2/docs")
from content import *
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x25, 0x63, 0xEB); INK = RGBColor(0x0F, 0x17, 0x2A); MUTE = RGBColor(0x64, 0x74, 0x8B)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
base = doc.styles["Normal"]; base.font.name = "Calibri"; base.font.size = Pt(10.5)
base.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
base.paragraph_format.space_after = Pt(6); base.paragraph_format.line_spacing = 1.12
for n, sz, col in [("Heading 1", 17, INK), ("Heading 2", 13, BLUE), ("Heading 3", 11, INK)]:
    s = doc.styles[n]; s.font.name = "Calibri"; s.font.size = Pt(sz); s.font.bold = True; s.font.color.rgb = col
    s.element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); s.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    s.paragraph_format.space_before = Pt(14 if n == "Heading 1" else 10); s.paragraph_format.space_after = Pt(5); s.paragraph_format.keep_with_next = True

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)
def H1(t): doc.add_heading(t, 1)
def H2(t): doc.add_heading(t, 2)
def H3(t): doc.add_heading(t, 3)
def P(t, bold_lead=None, italic=False):
    p = doc.add_paragraph()
    if bold_lead: r = p.add_run(bold_lead); r.bold = True
    r = p.add_run(t); r.italic = italic
    return p
def bl(items):
    for i in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(i, tuple): r = p.add_run(i[0]); r.bold = True; p.add_run(i[1])
        else: p.add_run(i)
        p.paragraph_format.space_after = Pt(2)
def callout(t, fill="DBEAFE"):
    tb = doc.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tb.cell(0, 0); shade(c, fill); c.paragraphs[0].add_run(t).font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
def table(head, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(head)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; shade(c, "2563EB"); r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if k % 2: shade(cells[i], "F1F5F9")
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.0
    if widths:
        t.autofit = False
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
        for i, w in enumerate(widths): t.columns[i].width = Inches(6.5 * w)
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(6.5 * w)
    for row in t.rows:
        cs = OxmlElement('w:cantSplit'); cs.set(qn('w:val'), 'true'); row._tr.get_or_add_trPr().append(cs)
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
def fig(path, cap, width=6.4):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width))
    c = doc.add_paragraph(); r = c.add_run(cap); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = MUTE; c.alignment = WD_ALIGN_PARAGRAPH.CENTER
def code(t):
    p = doc.add_paragraph(); r = p.add_run(t); r.font.name = "Consolas"; r.font.size = Pt(9); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.space_after = Pt(2)

# footer page numbers
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("ADAPT 2.0 documentation  |  page "); r.font.size = Pt(8.5); r.font.color.rgb = MUTE
r2 = fp.add_run(); r2.font.size = Pt(8.5)
for t, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
    if t: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
    else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
    r2._r.append(e)

dm = S["unnec_diff_mean"]; dci = S["unnec_diff_ci"]; A = M["adapt_selective"]; B = M["full_reschedule"]

# ================================================================ cover
for _ in range(5): doc.add_paragraph()
p = doc.add_paragraph(); r = p.add_run("ADAPT 2.0"); r.bold = True; r.font.size = Pt(38); r.font.color.rgb = INK
p = doc.add_paragraph(); r = p.add_run("AI-Powered Adaptive Student Learning and Support Platform"); r.font.size = Pt(16); r.font.color.rgb = BLUE
p = doc.add_paragraph(); r = p.add_run("Project documentation: requirements, design, algorithms, testing and results"); r.font.size = Pt(12); r.font.color.rgb = MUTE
doc.add_paragraph()
P(f"Document date: {TODAY}")
P("Project type: personal / portfolio project, built on the ADAPT adaptive planning engine (project proposal dated 10 October 2026)")
doc.add_paragraph()
callout("Status of this document. Stage 1, the C++17 planning engine, has been implemented and tested; every result in sections 13 and 14 was measured on that engine and can be reproduced from the project files. "
        "Stages 2 to 5 (notes library, OCR, AI explanations, quizzes, originality tools, Urdu and low-bandwidth support) are designed here but not built. No results are reported for them, and their test tables are left blank on purpose.", "FEF3C7")
doc.add_page_break()

H1("Contents")
for i, t in enumerate(["Abstract", "Introduction and problem statement", "Aim and objectives", "Scope and how ADAPT 2.0 differs from ADAPT 1.0", "Users and use cases", "Requirements",
                       "Workflow pipeline", "System architecture", "Engine design (Stage 1)", "ADAPT 2.0 modules (Stages 2 to 5, design only)", "Worked example", "Testing methodology", "Test results (Stage 1)",
                       "Interpretation, defect log and limits", "Test plans for Stages 2 to 5", "Roadmap and project plan", "Risks, ethics and limitations", "Conclusion", "References",
                       "Appendices: A reproduction guide, B sample data, C viva questions, D submission checklist, E glossary"], 1):
    P(f"{i}. {t}")
doc.add_page_break()

# ================================================================ 1 abstract
H1("1. Abstract")
P("ADAPT is an adaptive planning engine that repairs a student's study plan when reality changes, instead of rebuilding the whole plan. ADAPT 2.0 widens the project into a student learning and support platform: "
  "students upload notes, the system extracts and explains the material, generates practice questions, and turns weak topics into revision tasks that the C++ planner schedules, repairs after disruptions and re-estimates from feedback.")
P(f"Stage 1, the planner, is implemented in C++17 using only the standard library. It passed all {PASSED} planned scenario tests. On {S['n']} randomised scenarios it scheduled as many tasks as a full rescheduler "
  f"while moving {f2(dm)} fewer unaffected tasks per disruption ({f2(B['unnecessary'])} vs {f2(A['unnecessary'])}), and it produced no hard-constraint violations. "
  "The remaining stages are specified here with their own test plans and will be reported only after they are built and measured.")

# ================================================================ 2 intro
H1("2. Introduction and problem statement")
P("Most planners produce a schedule and stop. When a class is cancelled, a task takes twice as long, or an urgent assignment appears, the student has to rebuild the plan by hand. "
  "A tool that rebuilds everything from scratch can also move tasks that were fine, which makes the plan feel unstable and hard to trust.")
P("Students also lose time outside planning: notes are scattered across phones, folders and chat apps, difficult material is hard to understand, and many students have limited bandwidth, money or access to tutoring. ADAPT 2.0 treats these as one workflow rather than separate tools.")
P("Research question for Stage 1: can a C++ planning engine selectively repair a disrupted schedule, keep hard constraints, explain its changes and improve its duration estimates from feedback, while moving fewer unaffected tasks than a full rescheduler? ", "")
P("Design principle for the whole project: ADAPT should help students learn, not just produce answers for them.", italic=True)

H1("3. Aim and objectives")
bl(["Build a working C++ prototype that repairs study plans after disruptions (done for Stage 1).",
    "Represent tasks, priorities, deadlines, durations, available time and fixed commitments (done).",
    "Detect which tasks a disruption affects and keep the rest where they are when that is feasible (done).",
    "Score alternative plans with an explicit cost function and reject any plan that breaks a hard constraint (done).",
    "Explain every proposed change and let the user approve or reject it (done).",
    "Record actual outcomes and update duration estimates with a documented rule (done).",
    "Evaluate against a fixed-plan and a full-reschedule baseline (done on synthetic scenarios; user pilot pending).",
    "Add a notes library, OCR, AI explanations, quizzes and originality tools around the engine (designed, not built).",
    "Make the basic planner and notes library usable on slow connections and in simple English, with Urdu support where feasible (designed, not built)."])

# ================================================================ 4 scope
H1("4. Scope and how ADAPT 2.0 differs from ADAPT 1.0")
P("ADAPT 1.0 is a focused planning engine with a command-line interface. ADAPT 2.0 keeps that engine as its core and adds learning and support features around it. "
  "The two source documents disagreed on several points, so the following decisions were made and are used throughout this documentation.")
table(["Topic", "ADAPT 1.0 proposal", "ADAPT 2.0 idea", "Decision in this document"],
      [["Language", "C++17 only; 'no Python engine'", "C++ plus web, OCR and AI services", "C++17 owns scheduling, constraints and repair. Other languages may be used for the web interface, OCR and AI calls; they never make scheduling decisions."],
       ["Deadlines", "Hard boundary in some places, 'weighted lateness' in others", "Not specified", "Deadlines are hard. A task that cannot finish in time is reported as unscheduled. The lateness metric is replaced by unscheduled priority points."],
       ["Storage", "CSV / JSON; SQLite later", "SQLite initially", "CSV for the engine now; SQLite when the notes library is built."],
       ["Users and accounts", "Single user; no cloud accounts", "Account database", "Single local user for Stage 1. Accounts only if the project becomes a shared service, with the privacy plan in section 10.6."],
       ["Privacy", "Local storage, fictional data", "Uploads to OCR and AI services", "Users must be told when text leaves their device; uploads can be deleted; no use of private notes for training."],
       ["Timeline", "10 weeks, one semester", "Five stages, no dates", "Stage 1 is the semester deliverable. Stages 2 to 5 are future work (section 16)."],
       ["Machine learning", "Not in first version", "LLMs for explanations and quizzes", "No model is trained. Hosted or licensed models are called only from Stage 3 onward."]],
      [0.13, 0.22, 0.22, 0.43], 8.5)
P("Out of scope for Stage 1: notes upload, OCR, AI features, accounts, calendars, mobile apps, medical or psychological advice, automatic actions without approval, and any claim of perfect prediction.")

H1("5. Users and use cases")
P("Primary user: a university student managing several assignments, quizzes and exams. Secondary evaluators: a teacher reviewing the design and a small group of volunteer students testing usability.")
table(["ID", "Use case", "Stage"],
      [["UC-01", "Create a task: title, estimated duration, deadline, priority, splittable or not", "1"], ["UC-02", "Set availability: study windows and fixed commitments", "1"], ["UC-03", "Generate a baseline plan", "1"],
       ["UC-04", "Report a disruption: lost time, a new task or a changed duration", "1"], ["UC-05", "Repair the plan, preserving unaffected tasks", "1"], ["UC-06", "Compare alternatives and read the reasons for each change", "1"],
       ["UC-07", "Approve or reject a proposal; the decision is logged", "1"], ["UC-08", "Record actual duration so estimates improve", "1"],
       ["UC-09", "Upload notes (PDF, Word, photo) and see the extracted text", "2"], ["UC-10", "Organise and search notes by subject, semester and topic", "2"],
       ["UC-11", "Get a summary and simple explanation of a topic", "3"], ["UC-12", "Practise with flashcards and a quiz; see which topics need revision", "3"],
       ["UC-13", "Send quiz-driven revision tasks to the planner", "3"], ["UC-14", "Check writing for overlap with sources and get citation help", "4"]], [0.12, 0.76, 0.12])

# ================================================================ 6 requirements
H1("6. Requirements")
H2("6.1 Functional requirements")
table(["ID", "Requirement", "Stage", "Status"],
      [["FR-01", "Create, edit, list and delete tasks", "1", "Engine data model implemented; CLI menu pending"], ["FR-02", "Each task has unique id, title, duration, priority, deadline, splittable flag", "1", "Implemented and validated"],
       ["FR-03", "Store study windows and fixed commitments", "1", "Implemented (free-time intervals)"], ["FR-04", "Generate a baseline plan from valid input", "1", "Implemented"],
       ["FR-05", "Accept lost time, new tasks and changed durations", "1", "Implemented"], ["FR-06", "Identify affected tasks; preserve unaffected tasks where possible", "1", "Implemented"],
       ["FR-07", "Score candidate plans; explain the main reasons for changes", "1", "Implemented"], ["FR-08", "User approves or rejects a proposal", "1", "Implemented (session object)"],
       ["FR-09", "Record planned duration, actual duration and feedback", "1", "Estimator implemented; persistence pending"], ["FR-10", "Update estimates with a documented rule", "1", "Implemented (EMA, alpha 0.3)"],
       ["FR-11", "Export plans and test logs to CSV or text", "1", "Test logs and benchmark exported as CSV"],
       ["FR-12", "Upload PDF, Word and image files within a stated size limit", "2", "Planned"], ["FR-13", "Extract text; flag low-confidence OCR regions instead of guessing", "2", "Planned"],
       ["FR-14", "Organise notes by subject, semester, topic; full-text search", "2", "Planned"], ["FR-15", "Let the user delete uploads and derived data", "2", "Planned"],
       ["FR-16", "Summaries and simple-English explanations tied to source passages; Urdu where feasible", "3", "Planned"], ["FR-17", "Flashcards and quizzes generated from notes and reviewed against the source", "3", "Planned"],
       ["FR-18", "Create revision tasks from quiz results and pass them to the C++ engine", "3", "Planned"], ["FR-19", "Similarity check with source references; citation help", "4", "Planned"],
       ["FR-20", "AI-writing indicator shown as advisory only, with its limits explained", "4", "Planned"], ["FR-21", "Basic planner keeps working when AI services are unavailable", "3", "Planned"]],
      [0.09, 0.55, 0.07, 0.29], 8.5)
H2("6.2 Non-functional requirements")
table(["ID", "Requirement", "Evidence so far"],
      [["NFR-01", "Hard constraints are checked before a plan is shown as feasible", f"Independent checker; 0 violations in {S['rows']:,} method runs"],
       ["NFR-02", "Every moved or unscheduled task has a human-readable reason", "Implemented (see section 11)"],
       ["NFR-03", "Malformed input must not crash the program", "T6 and T10 passed"],
       ["NFR-04", "100 tasks scheduled and repaired in under 2 seconds (target)", f"Measured about {f2(S['scale'][2]['repair_adapt_us_median']/1000)} ms (median) on the test machine"],
       ["NFR-05", "Privacy: store only what planning needs; tell users when text leaves the device", "Local-only in Stage 1; policy for Stages 2 to 4 in section 10.6"],
       ["NFR-06", "Core engine compiles with a standard C++17 compiler", "g++ 13.3.0, no warnings with -Wall -Wextra"],
       ["NFR-07", "Modular design: model, planner, repair, scoring, storage, tests", "Engine split into logical modules (single header for now)"],
       ["NFR-08", "Low-bandwidth: pages and previews kept small", "Target to be set and measured in Stage 5"],
       ["NFR-09", "Cost transparency: free limits stated honestly; no promise of unlimited storage or processing", "Policy decision"]], [0.1, 0.5, 0.4], 8.5)

# ================================================================ 7 workflow
H1("7. Workflow pipeline")
P("The strongest idea in ADAPT 2.0 is that its features feed each other. A photo of notes becomes searchable text, then an explanation and a quiz; weak quiz topics become revision tasks; the planner fits those tasks around the student's real week and repairs the plan when something changes; actual times flow back into better estimates.")
fig("figs/fig1_pipeline.png", "Figure 1. ADAPT 2.0 workflow. Blue steps are implemented and tested (Stage 1); amber dashed steps are planned.")
table(STEPS_H, STEPS, [0.05, 0.13, 0.42, 0.18, 0.22], 8.5)

# ================================================================ 8 architecture
H1("8. System architecture")
P("The C++ engine handles planning and decision logic. Parsers, OCR, AI models and similarity tools handle language and documents. They communicate through the application backend, which means the planner can be developed, tested and used without any of them.")
fig("figs/fig2_architecture.png", "Figure 2. Proposed architecture. Blue: implemented in Stage 1. Amber dashed: planned.")
table(MODULES_H, MODULES, [0.22, 0.56, 0.22])
H3("Suggested technology")
table(["Component", "Technology", "Purpose"],
      [["Planning engine", "C++17", "Scheduling, constraints, repair, estimator"], ["Interface", "CLI now; HTML, CSS, JavaScript later", "Task entry, plan view, uploads, reader"], ["Backend", "A suitable web API", "Connects interface to engine and services"],
       ["Database", "CSV now; SQLite, then PostgreSQL if needed", "Tasks, plans, history, notes metadata"], ["File storage", "Local folder for prototype; object storage later", "PDFs and images"],
       ["Text extraction", "PDF / Word parsers and an OCR engine", "Make scanned notes searchable"], ["AI explanations", "Hosted model via API or a suitably licensed local model", "Summaries, explanations, quizzes"],
       ["Similarity", "Text comparison plus a legitimate source-checking service if available", "Possible matches with references"]], [0.2, 0.42, 0.38], 8.5)
fig("figs/fig4_loop.png", "Figure 3. The Plan - Observe - Diagnose - Repair - Learn loop.", 5.0)

# ================================================================ 9 engine
H1("9. Engine design (Stage 1)")
H2("9.1 Data model")
table(["Object", "Fields"],
      [["Task", "id, title, category, duration (minutes), priority 1 to 5, hard deadline (minute since Monday 00:00), splittable"], ["Interval", "start, end (free study time after fixed commitments are removed)"],
       ["Chunk", "task id, start, end (a splittable task may have several chunks)"], ["Plan", "chunks, unscheduled task ids, cost, explanation notes"], ["World", "free intervals and tasks"],
       ["Disruption", "lost window, new task, or changed duration"], ["Session", "approved plan, pending proposal, history log"]], [0.2, 0.8])
H2("9.2 Initial planning heuristic")
bl(["Validate every task (positive duration, priority 1 to 5, title, unique id, positive deadline).",
    "Sort by earliest deadline, then higher priority, then shorter duration.",
    "Place each task in the earliest free time that fits and ends before its deadline. A non-splittable task needs one continuous gap; a splittable task may use several.",
    "If a task cannot fit, list it as unscheduled with the reason. The engine never pretends it fits."])
H2("9.3 Selective repair")
bl(["Apply the disruption to the world (remove lost time, add a task, or change a duration).",
    "Mark affected tasks: those overlapping lost time, the task whose duration changed, or the new task.",
    "Candidate C1 (selective): freeze every unaffected task and re-place only the affected ones.",
    "If the disruption left a previously scheduled (or new) task unscheduled, also build C2 (expanded: later-deadline tasks may move) and C3 (full reschedule).",
    "Reject any candidate that fails the hard-constraint check; score the rest; pick the lowest cost.",
    "Write a reason for every moved or unscheduled task and show the proposal. The user approves or rejects; a rejection keeps the old plan and is logged."])
fig("figs/fig3_repair_flow.png", "Figure 4. Selective repair with approval.", 4.4)
H2("9.4 Hard constraints and cost")
P("Hard constraints, checked independently of the planner: no two chunks overlap; every chunk lies inside free time (so fixed commitments are never used); every chunk ends before its deadline; a scheduled task receives exactly its duration; a non-splittable task is never split.")
P(COSTFORM + ". Plans that violate a hard constraint are disqualified rather than penalised. The weights are initial design choices; section 13.4 shows the effect of changing the first one.")
H2("9.5 Duration estimator")
P("For each task category the engine keeps an estimate E. When the user records an actual duration A, it updates E_new = alpha x A + (1 - alpha) x E_old with alpha = 0.3. Missing or non-positive outcomes are ignored. A category with fewer than three observations is marked low-confidence so the interface can say so. Test T7 confirms E = 60 and A = 80 gives 66.0.")
H2("9.6 File format")
P("Tasks are loaded from CSV with columns id, title, duration_min, priority, deadline_min, splittable. Loading is all-or-nothing: any malformed line or duplicate id returns errors and no tasks, so a bad file can never half-replace good data (test T10).")

# ================================================================ 10 ADAPT 2.0 modules
H1("10. ADAPT 2.0 modules (Stages 2 to 5, design only)")
callout("Nothing in this section has been built or tested. It is the specification the later stages will be built and measured against.", "FEF3C7")
H2("10.1 Notes library and OCR (Stage 2)")
bl(["Accepted files: PDF, Word, common image formats, within a published size limit.", "Digital PDFs and Word files go through a parser; photographs and scans go through OCR.",
    "OCR output keeps a confidence value per region. Unclear regions are shown as unclear, not silently guessed, especially for handwriting.", "Notes are organised by subject, semester and topic and searched by full text.",
    "Users can delete a file and everything derived from it (text, index entries, generated material)."])
H2("10.2 Notes explainer (Stage 3)")
bl(["Short and detailed summaries, key terms and definitions, step-by-step explanations of hard concepts, in simple English; Urdu where quality can be checked by a fluent speaker.",
    "Every explanation links back to the passage of the user's notes it is based on, so it can be checked.", "The model is told to say when the notes do not contain the answer."])
H2("10.3 Quiz and revision rule (Stage 3)")
P("Quiz results are a learning indicator, not a measurement of understanding. A revision task is created only after at least 5 questions on a topic, using this initial rule (design values, to be tuned with real data):")
table(["Quiz score on topic", "Revision task", "Priority"], [["below 50%", "45 minutes", "5"], ["50% to below 75%", "30 minutes", "4"], ["75% or above", "15-minute spaced review", "2"]], [0.4, 0.4, 0.2])
P("The task's deadline is set before the exam date the student entered. The C++ engine then schedules it like any other task, and repairs it when plans change.")
H2("10.4 Originality tools (Stage 4)")
bl(["<b>Similarity check and citation help.</b> Highlights passages that overlap with available sources, shows the source where known, finds missing quotation marks or citations, and explains citation styles such as APA. It can only compare against sources it has access to, so a clean result is never proof of originality.".replace("<b>", "").replace("</b>", ""),
    "The feature is called an Originality Assistant: it helps students fix problems and understand sources. It does not promise to remove plagiarism or avoid detection.",
    "<b>AI-writing indicator.</b> Advisory only, shown with an explanation of the signals used and a warning that detectors make mistakes. ADAPT never labels anyone as cheating from a score.".replace("<b>", "").replace("</b>", ""),
    "Published research found that common AI-text detectors wrongly flag writing by non-native English speakers more often (Liang et al., 2023). Because many target users write English as a second language, the indicator must be tested for false positives on such text (test S4-4) before it is shown to anyone, and may be dropped if it cannot be made fair."])
H2("10.5 Accessibility and low-bandwidth support (Stage 5)")
bl(["Compressed previews and light pages; saved notes can be reopened without re-uploading.", "Simple-English explanations; Urdu support where it can be reviewed by a fluent speaker.",
    "The basic planner and notes library keep working if an external AI service is down.", "Free limits are stated honestly. Truly unlimited storage or processing cannot be promised."])
H2("10.6 Privacy and data handling")
bl(["Store only what the feature needs. Keep demonstration data fictional.", "Tell the user before any text leaves their device for OCR or AI processing.", "Provide deletion of uploads and derived data.",
    "Do not use private notes to train models. Do not collect grades, passwords or private messages in pilots."])

# ================================================================ 11 worked example
H1("11. Worked example")
P("Data are the five fictional tasks from the original proposal. Study windows are 16:00 to 22:00 Monday to Thursday, with a fixed commitment on Monday 19:00 to 20:00. The disruption removes Monday 16:00 to 18:00.")
table(SAMPLE_H, SAMPLE, [0.08, 0.28, 0.16, 0.12, 0.22, 0.14])
fig("figs/fig6_worked_example.png", "Figure 5. The same disruption handled by ADAPT and the two baselines (output of the real engine).")
table(EXAMPLE_H, EXAMPLE, [0.26, 0.2, 0.18, 0.16, 0.2])
P("A01 has a Monday 18:00 deadline, so once 16:00 to 18:00 is gone it cannot be done in time. ADAPT reports that instead of hiding it. It keeps A04 and A05 exactly where they were; the full reschedule moves both for no benefit. Explanations produced by the engine:")
bl(["A01 cannot be scheduled before its deadline (directly affected by the disruption)", "A02 moved to Mon 21:45 (directly affected by the disruption)", "A03 moved to Mon 18:00 (directly affected by the disruption)"])

# ================================================================ 12 methodology
H1("12. Testing methodology")
bl([("Scenario tests T1 to T10 ", "from the original proposal's test matrix, run on the demo data."),
    ("Randomised benchmark: ", f"{S['n']} scenarios = 3 load levels (8, 14, 20 tasks) x 3 disruption types x 100 seeds (generator seed 2026). 4-day horizon; study windows 09:00-12:00 and 16:00-22:00; about half the days have a 60-minute fixed commitment; task length 30-120 min; priority 1-5; 60% splittable. Lost windows and longer tasks are applied to a task that is actually scheduled."),
    ("Methods compared on identical inputs: ", "A fixed plan (affected tasks are dropped, nothing repaired), B full reschedule (same heuristic, rebuild everything), and ADAPT selective repair."),
    ("Metrics: ", "share of scenarios with all tasks scheduled; unscheduled priority points; previously scheduled tasks whose time changed; unaffected tasks moved; hard-constraint violations; runtime."),
    ("Scale test: ", "25 to 400 tasks over 40 days, median of 21 runs per size."),
    ("Environment: ", ENV + ". Timings are not from a student laptop.")])

# ================================================================ 13 results
H1("13. Test results (Stage 1)")
H2("13.1 Scenario tests T1 to T10")
table(TEST_H, TEST_ROWS, [0.06, 0.26, 0.56, 0.12], 8.5)
P(f"{PASSED} of {len(TESTS)} passed. These are functional checks on one demo dataset; they show the required behaviours exist, not how well the engine performs in general.")
H2("13.2 Benchmark results")
table(OVERALL[0], OVERALL[1], [0.2, 0.12, 0.14, 0.15, 0.12, 0.14, 0.13], 8.5)
fig("figs/fig7_benchmark.png", "Figure 6. Results by load level (900 scenarios). Whiskers are 95% confidence intervals.")
table(BYLOAD_H, BYLOAD, [0.14, 0.2, 0.22, 0.22, 0.22], 8.5)
fig("figs/fig8_by_disruption.png", "Figure 7. Unaffected tasks moved, by disruption type.", 5.0)
table(BYTYPE_H, BYTYPE, [0.4, 0.3, 0.3])
callout(f"Main finding. ADAPT completed as many scenarios as the full rescheduler ({pct(A['all_scheduled'])}) and left slightly fewer priority points unscheduled ({f2(A['unsched_weight'])} vs {f2(B['unsched_weight'])}), "
        f"while moving {f2(dm)} fewer unaffected tasks per scenario (95% CI +/- {f2(dci)}; {f2(B['unnecessary'])} vs {f2(A['unnecessary'])}). It moved fewer unaffected tasks in {S['adapt_fewer_moves_pct']:.0f}% of scenarios and the same number in {S['adapt_equal_moves_pct']:.0f}%; it never moved more. "
        f"Hard-constraint violations across all {S['rows']:,} method runs: {S['total_violations']}.")
H2("13.3 Selected strategy")
sh = S["strategy_share"]
P(f"With the default weights, ADAPT kept the purely selective plan in {sh.get('selective', 0)*100:.0f}% of scenarios, chose the expanded plan in {sh.get('expanded', 0)*100:.0f}% and fell back to a full reschedule in {sh.get('full', 0)*100:.1f}%.")
H2("13.4 Choosing the cost weight")
P("The weight on unscheduled priority was swept while everything else was held fixed. Lower weights make ADAPT more conservative: fewer unaffected tasks move, but more work is left unscheduled.")
fig("figs/fig9_weight_sweep.png", "Figure 8. Stability versus completeness as the weight changes.", 6.2)
table(SWEEP_H, SWEEP, [0.3, 0.2, 0.25, 0.25])
P("A weight near 20 keeps most of the completeness (52% vs 55% fully scheduled) while moving about 19% fewer unaffected tasks than the default. The right trade-off depends on how much students value stability, and should be decided with user feedback.")
H2("13.5 Runtime")
fig("figs/fig10_runtime.png", "Figure 9. Median repair time versus number of tasks.", 5.0)
table(SCALE_H, SCALE, [0.3, 0.35, 0.35])
P("The 100-task target of under 2 seconds is met by a wide margin (about 0.6 ms). ADAPT is slower than a plain rebuild because it evaluates several candidates, but stays within a few milliseconds even at 400 tasks.")
H2("13.6 Raw data sample")
P("First two benchmark scenarios (default weights). Full data: results/benchmark.csv (13,500 rows).")
table(RAW_H, RAW_T, [0.09, 0.06, 0.12, 0.14, 0.09, 0.1, 0.08, 0.08, 0.1, 0.06, 0.08], 7.5)

H1("14. Interpretation, defect log and limits")
H2("14.1 Defect found by testing")
P(f"The first benchmark run showed ADAPT moving {f2(S['v0_unnec_adapt'])} unaffected tasks per scenario, more than expected. The suspected cause was that ADAPT escalated to wider candidates even when the unscheduled tasks had already been unscheduled before the disruption. "
  "The trigger was changed to react only to newly unscheduled tasks, but the result barely changed (4.01 to 4.00). The real cause is the cost function: with 200 points per unscheduled priority point and 3 per unnecessary move, ADAPT will reshuffle many tasks to rescue one. "
  "That is a modelling choice, not a bug, which is why section 13.4 sweeps the weight. The first run is kept as results/benchmark_v0.1_defect.csv.")
H2("14.2 Limits of these results")
bl([("Synthetic scenarios. ", "They come from one random generator. Real student timetables may behave differently; the planned user pilot is the real test."),
    ("ADAPT includes the full reschedule as a fallback candidate. ", "So it cannot leave more priority points unscheduled than the full reschedule by construction. The meaningful comparison is the number of unaffected tasks moved."),
    ("Baseline A is deliberately weak. ", "It drops affected tasks without trying to repair them; it mainly shows the cost of doing nothing."),
    ("Timings come from a cloud container, ", "not a student's laptop."),
    ("Only Stage 1 is tested. ", "Nothing here says whether OCR, AI explanations, quizzes or originality tools will work well."),
    ("No novelty claim. ", "Adaptive scheduling and constraint optimisation are established research areas. This project claims a specific selective-repair design and measures it; it does not claim to be the first or the best.")])

# ================================================================ 15 stage test plans
H1("15. Test plans for Stages 2 to 5")
P("These tables are blank on purpose. The last column is filled only with measured results after each stage is built.")
for nm, (h, r) in [("Stage 2 - Notes library and OCR", S2), ("Stage 3 - AI explainer and quizzes", S3), ("Stage 4 - Originality tools", S4), ("Stage 5 - Accessibility and pilot", S5)]:
    H3(nm); table(h, r, [0.08, 0.34, 0.38, 0.2], 8.5)
P("Stage 5 pilot: invite 5 to 10 volunteer students, subject to instructor approval and university rules. Use fictional or non-sensitive tasks, and do not collect grades, passwords or private messages. Compare time to repair a plan manually versus with ADAPT, unnecessary moves, and 1 to 5 ratings for clarity, trust and ease of use. The sample is small and exploratory and cannot prove general effectiveness.")

# ================================================================ 16 roadmap
H1("16. Roadmap and project plan")
fig("figs/fig5_roadmap.png", "Figure 10. Staged plan. Only Stage 1 is built; week counts are indicative.", 5.8)
table(["Stage", "Deliverable", "Done when"],
      [["1 Core planner", "C++17 scheduler, repair, scoring, explanations, estimator, tests", "All T1-T10 pass; benchmark reported (done); CLI and user pilot pending"],
       ["2 Notes library", "Upload, extraction, OCR, organisation, search, delete", "S2 tests run and recorded"], ["3 AI explainer", "Summaries, explanations, flashcards, quizzes, quiz-to-task rule", "S3 tests run; human review of quiz quality recorded"],
       ["4 Originality tools", "Similarity, citation help, advisory AI-writing indicator", "S4 tests run, including false-positive check on non-native English text"], ["5 Support", "Low-bandwidth pages, Urdu, usability pilot, cost report", "S5 results recorded"]], [0.17, 0.45, 0.38], 8.5)
P("Remaining Stage 1 work: a menu-driven CLI around the engine, persistent history, and the usability pilot from the original proposal.")

H1("17. Risks, ethics and limitations")
table(["Risk", "Mitigation"],
      [["Duration estimates are wrong", "Show them as estimates; allow correction; show low confidence when data is sparse"], ["Schedules can be impossible", "Report infeasibility clearly; never claim all tasks fit"],
       ["Over-optimised plans ignore preferences", "User approval and manual override stay in the loop"], ["Private notes sent to external services", "Tell the user; allow deletion; no training on notes; local-only fallback for the planner"],
       ["AI explanations or quizzes are wrong", "Tie explanations to source passages; show when the notes lack an answer; human review of quiz samples"], ["AI-writing indicator is unfair", "Advisory only; test on non-native English text; drop it if it cannot be made fair"],
       ["OCR misreads handwriting", "Flag low-confidence regions; do not present guesses as fact"], ["Scope creep", "Stages are built and tested in order; later stages start only after earlier ones are measured"],
       ["Costs of AI and OCR", "State free limits honestly; keep the planner independent of paid services"]], [0.35, 0.65], 8.5)

H1("18. Conclusion")
P(f"ADAPT 2.0 keeps a working, tested planning engine at its centre and adds a learning workflow around it. The engine met every functional test and, on {S['n']} synthetic scenarios, matched a full rescheduler on completeness while moving fewer unaffected tasks and breaking no hard constraints. "
  "The rest of the platform is specified and has test plans, but is unbuilt. The project's value will be decided by measured results at each stage, not by the number of features.")

H1("19. References")
bl(["Liu, C. L. and Layland, J. W. (1973). Scheduling algorithms for multiprogramming in a hard-real-time environment. Journal of the ACM, 20(1), 46-61. (earliest-deadline-first scheduling)",
    "Liang, W., Yuksekgonul, M., Mao, Y., Wu, E. and Zou, J. (2023). GPT detectors are biased against non-native English writers. Patterns, 4(7). (AI-text detector bias)",
    "Google OR-Tools, C++ guide and constraint optimisation overview: https://developers.google.com/optimization (optional future solver; not used in Stage 1). Record the access date when you consult it.",
    "cppreference.com, C++ chrono and filesystem libraries: https://en.cppreference.com/w/cpp",
    "ADAPT project proposal and design documentation, 10 October 2026 (original proposal)."])

# ================================================================ appendices
doc.add_page_break()
H1("Appendix A - Reproduction guide")
P("Project layout (C++17, standard library only):")
for l in ["src/adapt.hpp          engine: scheduler, repair, baselines, estimator, session, CSV loader", "tests/run_tests.cpp    T1-T10, writes results/test_results.csv", "tests/bench.cpp        900-scenario benchmark, weight sweep and scale test",
          "tests/demo.cpp         worked example schedules", "docs/make_figs.py      charts, flowcharts and summary statistics", "results/               test_results.csv, benchmark.csv, benchmark_v0.1_defect.csv, scale.csv, demo_schedules.csv, summary_stats.json"]:
    code(l)
P("Commands:")
for l in ["g++ -std=c++17 -O2 -Wall -Wextra -o run_tests tests/run_tests.cpp && ./run_tests", "g++ -std=c++17 -O2 -o bench tests/bench.cpp && ./bench", "g++ -std=c++17 -O2 -o demo tests/demo.cpp && ./demo", "python3 docs/make_figs.py"]:
    code(l)
P("Times are minutes since Monday 00:00 (Tuesday 16:00 is 1440 + 960 = 2400). Benchmark timings vary slightly from run to run; counts and means of schedule metrics are deterministic for seed 2026.")

H1("Appendix B - Sample data")
table(SAMPLE_H, SAMPLE, [0.08, 0.28, 0.16, 0.12, 0.22, 0.14])
P("Fictional demonstration data. Do not use a real student's private timetable without consent. Study windows: 16:00-22:00 Monday to Thursday; fixed commitment Monday 19:00-20:00.")

H1("Appendix C - Viva questions (updated for ADAPT 2.0)")
QA = [("Why is the planning engine written in C++?", "C++ gives structured data models, efficient algorithms and clear modules. Scheduling, constraints and repair stay in C++; OCR and AI calls may use other tools through the backend."),
      ("Is ADAPT deep learning?", "No model is trained. The planner uses a heuristic, constraints and an exponential moving average. From Stage 3 it calls a hosted or licensed language model for explanations and quizzes."),
      ("What is the main contribution?", "Selective plan repair with explanations and feedback, measured against a fixed plan and a full reschedule, inside a workflow that connects notes, practice and planning."),
      ("How is it different from a normal planner?", "It identifies affected tasks, moves as few others as it can, explains each change, records the user's decision and learns durations."),
      ("What if no schedule is feasible?", "It reports which tasks cannot be scheduled and why. Deadlines are hard, so nothing is silently placed late."),
      ("How was it tested?", f"{len(TESTS)} scenario tests and {S['n']} randomised scenarios against two baselines, with an independent hard-constraint checker. A user pilot is still to do."),
      ("What did the results show?", f"Same completeness as a full reschedule with {f2(dm)} fewer unaffected tasks moved per scenario, no violations. The advantage depends on the cost weights, and the scenarios are synthetic."),
      ("Can you claim it is the first of its kind?", "No. Adaptive scheduling is an established field. The claim is limited to this design and its measured results."),
      ("What are the risks of the AI-writing indicator?", "Detectors make mistakes and can be unfair to non-native English writers. It is advisory only and must pass a false-positive test first."),
      ("What is the main limitation?", "Only Stage 1 exists and its tests are synthetic. Duration estimates and user input can be wrong, and some disruptions make it impossible to finish everything.")]
for i, (q, a) in enumerate(QA, 1):
    P(a, bold_lead=f"Q{i}. {q}  ")

H1("Appendix D - Submission checklist")
bl(["Confirm the project title and scope with the course instructor.", "Check university formatting, citation and plagiarism requirements.", "Replace 'planned' with measured results only after each stage is built.",
    "Add screenshots of the real program only after the CLI or interface exists.", "Record compiler version, OS, test dates, scenario data, metrics and defects (done for Stage 1).", "Verify each reference URL and record access dates.",
    "Keep source code, test data and this report under version control."])

H1("Appendix E - Glossary")
table(["Term", "Meaning"],
      [["Hard constraint", "A rule that must never be broken: no overlap, inside free time, before the deadline, full duration"], ["Selective repair", "Moving only the tasks a disruption affects, unless widening the search is needed"],
       ["Full reschedule", "Rebuilding the whole plan from scratch with the same heuristic"], ["EDF", "Earliest deadline first"], ["EMA", "Exponential moving average; the estimator's update rule"],
       ["Unscheduled priority points", "Sum of the priorities of tasks that could not be scheduled; lower is better"], ["Unnecessary move", "A task not affected by the disruption whose time changed anyway"],
       ["OCR", "Optical character recognition: reading text from images"]], [0.25, 0.75], 9)

doc.core_properties.title = "ADAPT 2.0 - Project documentation"; doc.core_properties.author = "ADAPT project"
doc.save("ADAPT_2.0_Project_Documentation.docx")
print("docx ok")
